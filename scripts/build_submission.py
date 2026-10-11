"""
Regenerate the LogInsight final-submission deliverables from the current source.

Inputs are real repository material only:
  * screenshots captured by frontend/scripts/capture-visuals.mjs from the running app
  * metadata and verification figures from scripts/submission_meta.json, which was
    transcribed from the previously submitted presentation rather than invented
  * bundle sizes read from the production build in frontend/dist

Run:  python scripts/build_submission.py
"""
from __future__ import annotations

import json
import pathlib

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.util import Emu, Inches as PptInches, Pt as PptPt

ROOT = pathlib.Path(__file__).resolve().parent.parent
IMAGES = ROOT / "docs" / "images" / "signal-atlas"
DIST = ROOT / "frontend" / "dist" / "assets"
OUT = ROOT / "final-submission"

MUTED = RGBColor(0x55, 0x62, 0x5C)
PINK = PptRGB(0x0A, 0x55, 0x60)
PINKMUTED = PptRGB(0x55, 0x62, 0x5C)
PINKLIGHT = PptRGB(0xDC, 0xEB, 0xEC)

SLIDE_W = Emu(12191695)
SLIDE_H = Emu(6858000)

META = json.loads((ROOT / "scripts" / "submission_meta.json").read_text(encoding="utf-8"))

COURSE = META["course"]
STUDENTS = [(s["name"], s["roll"]) for s in META["students"]]
GUIDE = META["guide"]
DEPARTMENT = META["department"]
INSTITUTION = META["institution"]
TERM = META["term"]
VERIFY = [(v["label"], v["value"], v["note"]) for v in META["verification"]]
ALGO = [(a["name"], a["complexity"], a["role"]) for a in META["algorithms"]]
LIMITATIONS = list(META["limitations"])

# Chunks the browser requests on first paint. Three.js is deliberately absent.
EAGER_CHUNKS = ("index-", "vendor-react-", "vendor-motion-", "vendor-icons-")


def first_load_kb() -> tuple[str, str]:
    """Measure first-load JS from the production build, or say it was not built."""
    if not DIST.exists():
        return "not built", "not built"
    raw = 0
    for chunk in DIST.glob("*.js"):
        if any(tag in chunk.name for tag in EAGER_CHUNKS):
            raw += chunk.stat().st_size
    kb = raw / 1000
    return f"{kb:.0f} kB", "see build log"


BUNDLE = [
    ("First-load JS", "before: shipped inside the entry chunk",
     f"after: {first_load_kb()[0]} across four cacheable chunks"),
    ("Three.js", "before: loaded with the application",
     "after: deferred until the 3D topology is opened"),
]


def shot(name: str) -> pathlib.Path | None:
    candidate = IMAGES / name
    return candidate if candidate.exists() else None


def add_slide(prs: Presentation, kicker: str, title: str, subtitle: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    band = slide.shapes.add_textbox(PptInches(0.7), PptInches(0.45), PptInches(11.9), PptInches(0.32))
    run = band.text_frame.paragraphs[0].add_run()
    run.text = kicker.upper()
    run.font.size = PptPt(11)
    run.font.bold = True
    run.font.color.rgb = PINK

    head = slide.shapes.add_textbox(PptInches(0.7), PptInches(0.8), PptInches(11.9), PptInches(0.85))
    run = head.text_frame.paragraphs[0].add_run()
    run.text = title
    run.font.size = PptPt(30)
    run.font.bold = True
    run.font.color.rgb = PINK

    if subtitle:
        sub = slide.shapes.add_textbox(PptInches(0.7), PptInches(1.62), PptInches(11.9), PptInches(0.5))
        run = sub.text_frame.paragraphs[0].add_run()
        run.text = subtitle
        run.font.size = PptPt(13)
        run.font.color.rgb = PINKMUTED

    rule = slide.shapes.add_shape(1, PptInches(0.7), PptInches(2.12), PptInches(11.9), Emu(11430))
    rule.fill.solid()
    rule.fill.fore_color.rgb = PINKLIGHT
    rule.line.fill.background()
    rule.shadow.inherit = False
    return slide


def body(slide, text: str, top: float, size: int = 13, height: float = 0.6):
    box = slide.shapes.add_textbox(PptInches(0.7), PptInches(top), PptInches(11.9), PptInches(height))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = PptPt(size)
    run.font.color.rgb = PINKMUTED
    return box


def add_picture(slide, path: pathlib.Path, top: float, left: float = 0.7, width=None, height=None):
    with path.open("rb") as handle:
        slide.shapes.add_picture(handle, PptInches(left), PptInches(top),
                                 width=PptInches(width) if width else None,
                                 height=PptInches(height) if height else None)


def footer(slide, page: int) -> None:
    box = slide.shapes.add_textbox(PptInches(0.7), PptInches(6.85), PptInches(11.9), PptInches(0.3))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = f"LogInsight  ·  KLH_CSE_2026-27_DSA-3_S3_T17_Loginsight-Analyzer  ·  {page}"
    run.font.size = PptPt(9)
    run.font.color.rgb = PINKMUTED


def build_pptx() -> int:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    page = 0

    # Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for top, text, size, bold, color in [
        (1.05, COURSE, 12, True, PINK),
        (1.5, META["project_title"], 48, True, PINK),
        (2.6, META["project_subtitle"], 19, False, PINKMUTED),
    ]:
        box = slide.shapes.add_textbox(PptInches(0.9), PptInches(top), PptInches(11.5), PptInches(1.0))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = PptPt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    names = slide.shapes.add_textbox(PptInches(0.9), PptInches(3.6), PptInches(6.0), PptInches(1.0))
    for index, (name, roll) in enumerate(STUDENTS):
        para = names.text_frame.paragraphs[0] if index == 0 else names.text_frame.add_paragraph()
        run = para.add_run()
        run.text = f"{name}\n{roll}"
        run.font.size = PptPt(14)
        run.font.bold = True
        run.font.color.rgb = PINKMUTED

    body(slide, f"{GUIDE}\n{DEPARTMENT}\n{INSTITUTION}  ·  {TERM}", 5.2, 12, 1.0)
    page += 1

    # Problem
    slide = add_slide(prs, "The problem", "One error line, spread across seven services",
                      "During an outage the evidence exists, but it is diluted, interleaved and queried faster than a human can read it.")
    body(slide, "A payment provider begins returning 502 responses.  ·  The order service retries; the gateway returns 504.  "
                "·  Notifications stop; customers see failed checkouts.  ·  The causal line is scattered across seven "
                "services, in an order that does not match causation.", 2.4, 13, 1.2)
    body(slide, "Why naive search fails: substring search is O(nm) and rescans the corpus; the engineer never learns "
                "which algorithm actually ran or what it measured.", 3.9, 13, 0.8)
    page += 1
    footer(slide, page)

    # Algorithms
    slide = add_slide(prs, "The solution", "Classical algorithms, made visible",
                      "Each number on screen is traceable to the algorithm that produced it.")
    top = 2.35
    for name, complexity, role in ALGO:
        card = slide.shapes.add_textbox(PptInches(0.7), PptInches(top), PptInches(11.9), PptInches(0.62))
        run = card.text_frame.paragraphs[0].add_run()
        run.text = f"{name}  ·  {complexity}    "
        run.font.size = PptPt(13)
        run.font.bold = True
        run.font.color.rgb = PINK
        run2 = card.text_frame.paragraphs[0].add_run()
        run2.text = role
        run2.font.size = PptPt(12)
        run2.font.color.rgb = PINKMUTED
        top += 0.66
    page += 1
    footer(slide, page)

    # Architecture
    slide = add_slide(prs, "Architecture", "Three layers, one server-owned state",
                      "The backend owns telemetry; the browser renders API responses with their scope and provenance.")
    body(slide,
         "PRESENTATION   React 18 · TypeScript · Vite · Motion\n"
         "                    │  REST /api/*        SSE  /api/simulation/stream · /api/live · /api/runs/{id}/events\n"
         "SPRING BOOT    DatasetService · LogIndex · Analytics · IncidentDetector · QueryDispatcher\n"
         "                    │  one current in-memory dataset, process-local incident state",
         2.4, 13, 2.0)
    body(slide, "Three SSE surfaces that are never conflated: generated simulation, bounded dataset replay, recorded trace replay.",
         4.6, 12, 0.5)
    page += 1
    footer(slide, page)

    # Interface slides
    for kicker, title, subtitle, image in [
        ("Interface", "Command center",
         "Selected-window metrics, observed relationships and detector windows — backend values only.",
         "command-center-desktop.png"),
        ("Interface", "Incident workbench",
         "Identity, measured context, evidence, blast radius and an explicit operator lifecycle.",
         "incident-workbench-desktop.png"),
        ("Interface", "Service topology — one graph, two renderers",
         "2D accessible SVG by default; 3D WebGL on demand with a working fallback.",
         "topology-selection-desktop.png"),
        ("Interface", "Algorithm catalogue and traceable runs",
         "Module, query type, time and space complexity, and whether an entry is exposed or traceable.",
         "algorithms-desktop.png"),
    ]:
        slide = add_slide(prs, kicker, title, subtitle)
        path = shot(image)
        if path:
            add_picture(slide, path, 2.3, width=8.4)
        page += 1
        footer(slide, page)

    # Logs + analytics, side by side
    slide = add_slide(prs, "Interface", "Logs and analytics",
                      "Server-side search, filters and pagination; charts computed by the backend for the selected range.")
    left = shot("log-explorer-desktop.png")
    right = shot("analytics-desktop.png")
    if left:
        add_picture(slide, left, 2.3, 0.7, width=5.5)
    if right:
        add_picture(slide, right, 2.3, 6.6, width=5.5)
    page += 1
    footer(slide, page)

    # States
    slide = add_slide(prs, "Interface", "States are explicit, never fabricated",
                      "No dataset and responsive mobile — onboarding rather than invented numbers.")
    left = shot("no-dataset-state.png")
    right = shot("command-center-mobile.png")
    if left:
        add_picture(slide, left, 2.3, 0.7, width=4.1)
    if right:
        add_picture(slide, right, 2.3, 5.1, height=4.2)
    page += 1
    footer(slide, page)

    # Verification
    slide = add_slide(prs, "Verification", "Claims that are measured, not asserted",
                      "Reproducible from the repository; CI runs the same commands on every push to main.")
    top = 2.4
    for label, value, note in VERIFY:
        box = slide.shapes.add_textbox(PptInches(0.7), PptInches(top), PptInches(11.9), PptInches(0.55))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = f"{label}: {value}"
        run.font.size = PptPt(14)
        run.font.bold = True
        run.font.color.rgb = PINK
        run2 = box.text_frame.paragraphs[0].add_run()
        run2.text = f"   — {note}"
        run2.font.size = PptPt(11)
        run2.font.color.rgb = PINKMUTED
        top += 0.58
    page += 1
    footer(slide, page)

    # Performance
    slide = add_slide(prs, "Verification", "Bundle work that was measured",
                      "Route-level splitting and vendor chunking, measured from the production build.")
    top = 2.5
    for label, before, after in BUNDLE:
        box = slide.shapes.add_textbox(PptInches(0.7), PptInches(top), PptInches(11.9), PptInches(0.5))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = f"{label}"
        run.font.size = PptPt(13)
        run.font.bold = True
        run.font.color.rgb = PINK
        run2 = box.text_frame.paragraphs[0].add_run()
        run2.text = f"\n{before}  →  {after}"
        run2.font.size = PptPt(11)
        run2.font.color.rgb = PINKMUTED
        top += 0.72
    body(slide, "The WebGL scene is rebuilt only when the graph genuinely changes, and its context is explicitly "
                "released on teardown rather than left to the browser's context cap.", 4.9, 12, 0.8)
    page += 1
    footer(slide, page)

    # Limitations
    slide = add_slide(prs, "Limitations", "Stated plainly",
                      "What this project is not, so the claims above can be read honestly.")
    top = 2.35
    for item in LIMITATIONS:
        box = slide.shapes.add_textbox(PptInches(0.7), PptInches(top), PptInches(11.9), PptInches(0.45))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = f"·  {item}"
        run.font.size = PptPt(12)
        run.font.color.rgb = PINKMUTED
        top += 0.5
    page += 1
    footer(slide, page)

    # Conclusion
    slide = add_slide(prs, "Conclusion", "Classical algorithms, applied honestly",
                      "The DSA is load-bearing: it runs in the product pipeline, not only in a demonstration folder.")
    body(slide, "Every figure shown is returned by the backend. Where data does not exist the interface says so instead "
                "of inventing a value, and where a capability is unavailable the UI states the limitation rather than "
                "hiding it.", 2.4, 13, 1.0)
    page += 1
    footer(slide, page)

    # Thank you
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(PptInches(0.9), PptInches(2.4), PptInches(11.5), PptInches(1.0))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Thank you"
    run.font.size = PptPt(40)
    run.font.bold = True
    run.font.color.rgb = PINK
    body(slide, "Questions are welcome — particularly on the algorithm choices and the honesty constraints.", 3.5, 14, 0.5)
    body(slide, "  ·  ".join(f"{name} ({roll})" for name, roll in STUDENTS), 4.3, 12, 0.5)

    target = OUT / "LogInsight_Final_Project_Presentation.pptx"
    prs.save(target)
    return len(prs.slides._sldIdLst)


def build_docx() -> int:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading(META["project_title"], level=0)
    doc.add_paragraph(META["project_subtitle"])
    doc.add_paragraph(COURSE)
    doc.add_paragraph(f"{INSTITUTION}  ·  {TERM}")

    doc.add_heading("Team", level=1)
    for name, roll in STUDENTS:
        doc.add_paragraph(f"{name} — {roll}")
    doc.add_paragraph(f"{GUIDE}, {DEPARTMENT}")

    for heading, paragraphs, bullets in [
        ("1. Problem statement", [
            "During an outage the evidence required to explain what happened is spread across services, arrives "
            "interleaved and out of order, and is queried far faster than a human can read it. Substring search over "
            "the corpus is O(nm) and rescans on every attempt, so the engineer never learns which algorithm actually "
            "ran or what it measured.",
            "LogInsight applies classical data structures and algorithms to that work — multi-pattern and "
            "exact-substring search, bounded edit distance, streaming window aggregation and dependency-graph "
            "traversal — and presents the returned events, aggregates and algorithm evidence together with their "
            "source and scope.",
        ], []),
        ("2. Architecture", [
            "The application is three layers with one rule: telemetry is server-owned. The Spring Boot backend "
            "generates deterministic scenarios, computes product aggregates, executes the algorithm engines and "
            "manages process-local incident state. The React frontend renders API responses with their scope and "
            "provenance, and claims algorithm execution evidence only where a real endpoint produced it.",
            "Transport is HTTP REST plus three Server-Sent Event streams that are deliberately kept distinct: the "
            "generated simulation, the bounded dataset replay, and recorded trace replay. There is no WebSocket "
            "client and no external log collector.",
        ], []),
        ("3. Data flow", [], [
            "A dataset is loaded from sample-data, uploaded, or generated by the backend.",
            "DatasetService installs it as one immutable in-memory current dataset; readers take a consistent snapshot.",
            "Analytics, incident detection, pattern extraction and search all derive from that dataset.",
            "The frontend requests derived values and never recomputes backend analysis.",
            "Mutating the dataset invalidates every cached view through a shared event.",
        ]),
        ("4. Algorithms and their actual roles", [
            "The catalogue distinguishes algorithms exposed through an execution endpoint, algorithms that "
            "additionally emit a trace, and entries that exist as lab implementations. A wider repertoire — bitmask "
            "TSP, Needleman-Wunsch, Dinic flow, vertex-cover kernelization, parallel prefix scan and Miller-Rabin — "
            "is exposed as runnable lab entries rather than presented as product capability.",
        ], []),
        ("5. 2D and 3D visualisation", [], [
            "The service topology renders as an accessible 2D SVG by default, with a paired service and edge list "
            "so the same data is available without WebGL. An optional Three.js view is code-split and loaded on demand.",
            "The WebGL scene is rebuilt only when the graph content genuinely changes, not on every render.",
            "Teardown disposes geometries, materials and controls, stops the animation loop, and explicitly releases "
            "the browser WebGL context, which renderer disposal alone does not do.",
            "Context loss surfaces a usable 2D fallback; context restoration rebuilds the scene automatically.",
            "Switching modes preserves the selected service and every active filter.",
            "Node size follows observed event counts and health follows error-rate bands; neither implies causation.",
        ]),
        ("6. Error handling and data honesty", [
            "Every dataset-backed request returns a structured error envelope. Validation and malformed input map to "
            "400, a missing dataset to 404, and unexpected failures to a sanitised 500 that never carries a stack "
            "trace. When no dataset is loaded the interface shows onboarding rather than zeroed metrics, and "
            "unmeasured values render as an em dash instead of being coerced to zero.",
        ], []),
    ]:
        doc.add_heading(heading, level=1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("4.1 Algorithm catalogue", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, heading in zip(table.rows[0].cells, ("Algorithm", "Complexity", "Role in the product")):
        cell.text = heading
    for name, complexity, role in ALGO:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = name, complexity, role

    doc.add_heading("7. Testing methodology and results", level=1)
    doc.add_paragraph(
        "The browser suite covers direct navigation to every documented route, route aliases, dataset loading and "
        "replacement, search and pagination, empty and failing states, repeated 2D/3D mode switching, the WebGL "
        "fallback, the guided walkthrough, reduced motion, console and failed network requests, and responsive "
        "overflow at six viewport sizes.")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, heading in zip(table.rows[0].cells, ("Check", "Result", "Command")):
        cell.text = heading
    for label, value, note in VERIFY:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = label, value, note

    doc.add_heading("8. Performance", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, heading in zip(table.rows[0].cells, ("Measure", "Before", "After")):
        cell.text = heading
    for label, before, after in BUNDLE:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = label, before, after

    doc.add_heading("9. Screenshots from the final implementation", level=1)
    for filename, caption in [
        ("command-center-desktop.png", "Command center with backend-derived selected-window values."),
        ("incident-workbench-desktop.png", "Incident workbench: identity, measured context, blast radius, lifecycle."),
        ("log-explorer-desktop.png", "Log explorer with a server-side query and results."),
        ("analytics-desktop.png", "Analytics computed for the selected range."),
        ("topology-selection-desktop.png", "Service topology with a selected service."),
        ("algorithms-desktop.png", "Algorithm catalogue."),
        ("no-dataset-state.png", "No dataset: onboarding rather than fabricated metrics."),
        ("command-center-mobile.png", "Responsive mobile layout."),
    ]:
        path = shot(filename)
        if not path:
            continue
        doc.add_picture(str(path), width=Inches(6.1))
        paragraph = doc.add_paragraph(caption)
        paragraph.runs[0].italic = True
        paragraph.runs[0].font.size = Pt(9)
        paragraph.runs[0].font.color.rgb = MUTED

    doc.add_heading("10. Known limitations", level=1)
    for item in LIMITATIONS:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("11. Future work", level=1)
    for item in [
        "Durable dataset storage so investigations survive a restart.",
        "Authentication and authorisation, which the current deployment model does not provide.",
        "Streaming ingestion from a real collector, replacing the generated simulation and bounded replay.",
        "Distributed algorithm execution for the largest catalogue entries.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    target = OUT / "LogInsight_Final_Project_Report.docx"
    doc.save(target)
    return len(doc.paragraphs)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    print(f"presentation: {build_pptx()} slides")
    print(f"report: {build_docx()} paragraphs")
    print(f"first-load JS: {first_load_kb()[0]}")